"""Run with: python manage.py test tests.browser_navigation (requires Playwright and Chrome)."""
from types import SimpleNamespace
from concurrent.futures import ThreadPoolExecutor
from tempfile import TemporaryDirectory
from datetime import date
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.contrib.staticfiles.testing import StaticLiveServerTestCase
from playwright.sync_api import sync_playwright, expect

from proxies.models import ProxyConfig


class NavigationBrowserTests(StaticLiveServerTestCase):
    def test_traffic_filter_and_excel_without_navigation(self):
        import json
        from io import BytesIO
        from pathlib import Path
        from openpyxl import load_workbook

        get_user_model().objects.create_superuser("traffic-admin", password="test-password")
        with TemporaryDirectory() as logs, self.settings(NGINX_ACCESS_LOG_DIR=logs), sync_playwright() as playwright:
            Path(logs, "test.access.log").write_text("\n".join(json.dumps({
                "destination_fqdn": domain, "time": "2026-09-14T12:00:00Z", "request": "GET /",
            }) for domain in ["one.example.org", "two.example.org"]), encoding="utf-8")
            browser = playwright.chromium.launch(channel="chrome", headless=True)
            page = browser.new_page()
            page.goto(self.live_server_url + "/login/?next=/audit/")
            page.locator('[name="username"]').fill("traffic-admin")
            page.locator('[name="password"]').fill("test-password")
            page.get_by_role("button", name="Sign in", exact=True).click()
            expect(page.locator("#traffic-rows")).to_contain_text("two.example.org")
            page.evaluate("window.documentToken = 'unchanged'")
            page.locator("#traffic-fqdn").fill("one.example.org")
            page.get_by_role("button", name="Filter", exact=True).click()
            expect(page.locator("#traffic-rows")).to_contain_text("one.example.org")
            expect(page.locator("#traffic-rows")).not_to_contain_text("two.example.org")
            with page.expect_response(lambda response: "/audit/rows/?fqdn=one.example.org" in response.url):
                page.wait_for_timeout(5500)
            expect(page.locator("#traffic-rows")).not_to_contain_text("two.example.org")
            for label, count in [("Export current view to Excel", 2), ("Export all to Excel", 3)]:
                with page.expect_download() as download:
                    page.get_by_role("link", name=label, exact=True).click()
                workbook = load_workbook(BytesIO(Path(download.value.path()).read_bytes()))
                self.assertEqual(workbook.active.max_row, count)
                workbook.close()
            page.get_by_role("link", name="Show all", exact=True).click()
            expect(page.locator("#traffic-rows")).to_contain_text("two.example.org")
            self.assertEqual(page.evaluate("window.documentToken"), "unchanged")
            browser.close()

    def test_workflows_keep_the_same_document(self):
        get_user_model().objects.create_superuser("browser-admin", password="test-password")
        with TemporaryDirectory() as media, self.settings(MEDIA_ROOT=media), patch("proxies.forms.certificate_valid_until", return_value=date(2030, 1, 1)), patch("proxies.forms.resolve_public_ip", return_value="8.8.8.8"), patch(
            "proxies.views.apply_proxy", return_value=SimpleNamespace(ok=True, message="Applied")
        ), sync_playwright() as playwright:
            browser = playwright.chromium.launch(channel="chrome", headless=True)
            page = browser.new_page()
            errors = []
            page.on("pageerror", lambda error: errors.append(str(error)))
            page.goto(self.live_server_url + "/login/")
            page.evaluate("window.documentToken = 'unchanged'")
            page.locator('[name="username"]').fill("browser-admin")
            page.locator('[name="password"]').fill("test-password")
            page.get_by_role("button", name="Sign in", exact=True).click()
            expect(page.locator(".stats")).to_be_visible()
            page.get_by_role("link", name="+ Add Reverse Proxy").click()
            page.locator('[name="domain_name"]').fill("not-a-domain")
            page.locator('[name="backend_private_ip"]').fill("10.0.0.5")
            page.locator('[name="backend_port"]').fill("8080")
            page.locator('[name="incoming_protocol"]').select_option("http")
            page.get_by_role("button", name="Save configuration").click()
            expect(page.locator(".errorlist")).to_be_visible()
            page.locator('[name="domain_name"]').fill("browser.example.org")
            page.get_by_role("button", name="Save configuration").click()
            expect(page.locator("#proxies")).to_contain_text("browser.example.org")
            page.locator("#proxies").get_by_role("switch").click()
            expect(page.locator("#proxies")).to_contain_text("Disabled")
            page.locator('#proxies button[title="Apply"]').click()
            expect(page.locator(".messages")).to_contain_text("Applied")
            page.locator('.sidebar a[href="/proxies/"]').click()
            expect(page.locator("h1")).to_have_text("Reverse Proxies")
            with page.expect_download() as download:
                page.get_by_role("link", name="Export all as PDF").click()
            self.assertEqual(download.value.suggested_filename, "reverse-proxies.pdf")
            page.locator('.sidebar a[href="/domains/"]').click()
            expect(page.locator("h1")).to_have_text("Domains & IPs")
            page.go_back()
            expect(page.locator("h1")).to_have_text("Reverse Proxies")
            page.go_forward()
            expect(page.locator("h1")).to_have_text("Domains & IPs")
            page.locator('.sidebar a[href="/servers/"]').click()
            expect(page.locator("#monitor-status")).to_contain_text("Updated")
            page.locator('.sidebar a[href="/audit/"]').click()
            expect(page.locator("#traffic-refresh-status")).to_be_empty()
            page.locator('.sidebar a[href="/servers/"]').click()
            expect(page.locator("#monitor-status")).to_contain_text("Updated")
            page.locator('.sidebar a[href="/certificates/"]').click()
            page.locator('[name="name"]').fill("Keep this unfinished upload")
            # Background inventory updates must not reset the upload form.
            page.wait_for_timeout(5500)
            expect(page.locator('[name="name"]')).to_have_value("Keep this unfinished upload")
            page.locator('[name="certificate"]').set_input_files({
                "name": "test.pem", "mimeType": "application/x-pem-file",
                "buffer": b"-----BEGIN CERTIFICATE-----\ntest\n-----END CERTIFICATE-----",
            })
            page.locator('[name="private_key"]').set_input_files({
                "name": "test.key", "mimeType": "application/x-pem-file",
                "buffer": b"-----BEGIN PRIVATE KEY-----\ntest\n-----END PRIVATE KEY-----",
            })
            page.get_by_role("button", name="Upload securely").click()
            expect(page.locator(".certificate-row")).to_contain_text("Keep this unfinished upload")
            page.once("dialog", lambda dialog: dialog.dismiss())
            page.locator(".certificate-row button").click()
            expect(page.locator(".certificate-row")).to_be_visible()
            page.once("dialog", lambda dialog: dialog.accept())
            page.locator(".certificate-row button").click()
            expect(page.locator(".certificate-row")).to_have_count(0)

            page.locator('.sidebar a[href="/proxies/"]').click()
            expect(page.locator("#proxies")).to_contain_text("browser.example.org")
            # Changes from another session appear without a document navigation.
            with ThreadPoolExecutor(max_workers=1) as pool:
                pool.submit(lambda: ProxyConfig.objects.filter(domain_name="browser.example.org").update(backend_port=9090)).result()
            expect(page.locator("#proxies")).to_contain_text("9090", timeout=10000)
            page.locator('#proxies a[title="Delete"]').click()
            expect(page.locator("h1")).to_have_text("Delete proxy")
            page.get_by_role("button", name="Delete proxy", exact=True).click()
            expect(page.locator("#proxies")).not_to_contain_text("browser.example.org")
            page.locator(".account-dropdown summary").click()
            page.get_by_role("button", name="Logout", exact=True).click()
            expect(page.locator(".login-panel")).to_be_visible()
            self.assertEqual(page.evaluate("window.documentToken"), "unchanged")
            self.assertEqual(errors, [])
            browser.close()

    def test_search_and_switch_on_loaded_navigated_and_live_pages(self):
        user = get_user_model().objects.create_superuser("search-admin", password="test-password")
        domains = ["alpha.find.example.org", "zulu.find.example.org"] + [
            f"middle-{index:02d}.example.org" for index in range(14)
        ]
        ProxyConfig.objects.bulk_create([
            ProxyConfig(domain_name=domain, backend_private_ip="10.0.0.5", backend_port=8080,
                        created_by=user, updated_by=user)
            for domain in domains
        ])
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(channel="chrome", headless=True)
            page = browser.new_page()
            errors = []
            page.on("pageerror", lambda error: errors.append(str(error)))
            page.goto(self.live_server_url + "/login/?next=/proxies/")
            page.locator('[name="username"]').fill("search-admin")
            page.locator('[name="password"]').fill("test-password")
            page.get_by_role("button", name="Sign in", exact=True).click()
            expect(page.locator("#proxies")).to_contain_text("alpha.find.example.org")
            # Exercise initialization on a full authenticated load before testing app navigation.
            page.reload()
            page.evaluate("window.documentToken = 'unchanged'")
            first = page.locator("#proxies tbody tr").filter(has_text="alpha.find.example.org")
            last = page.locator("#proxies tbody tr").filter(has_text="zulu.find.example.org")
            toggle = first.get_by_role("switch")
            expect(toggle).to_have_attribute("aria-checked", "true")
            expect(toggle).to_have_css("background-color", "rgb(22, 163, 74)")
            expect(toggle.locator(".proxy-switch-thumb")).to_have_css("transform", "matrix(1, 0, 0, 1, 20, 0)")
            toggle.click()
            expect(toggle).to_have_attribute("aria-checked", "false")
            expect(toggle).to_have_css("background-color", "rgb(220, 38, 38)")
            expect(toggle.locator(".proxy-switch-thumb")).to_have_css("transform", "none")
            expect(first).to_contain_text("Disabled")
            toggle.click()
            expect(toggle).to_have_attribute("aria-checked", "true")
            expect(first).to_contain_text("Online")

            search = page.get_by_role("searchbox", name="Search this page", exact=True)
            status = page.locator("#search-status")
            search.fill("find.example.org")
            expect(page.locator("#proxies .search-match")).to_have_count(2)
            expect(page.locator("#proxies .search-blink")).to_have_count(2)
            expect(status).to_contain_text("2 matching items")
            # Matching rows are far apart, so Enter must visibly move between them.
            search.press("Enter")
            expect(first).to_be_in_viewport()
            search.press("Enter")
            expect(last).to_be_in_viewport()
            expect(first).not_to_be_in_viewport()
            search.press("Escape")
            expect(search).to_have_value("")
            expect(status).to_be_hidden()
            expect(page.locator(".search-match, .search-blink")).to_have_count(0)

            page.locator('.sidebar a[href="/domains/"]').click()
            expect(page.locator("h1")).to_have_text("Domains & IPs")
            search.fill("alpha.find.example.org")
            expect(page.locator(".search-match")).to_have_count(1)
            expect(status).to_contain_text("1 matching item")
            page.locator('.sidebar a[href="/proxies/"]').click()
            expect(page.locator("h1")).to_have_text("Reverse Proxies")
            search.fill("fresh.example.org")
            expect(status).to_contain_text("No matching items")
            # An edit from another session replaces table rows while the query remains active.
            with ThreadPoolExecutor(max_workers=1) as pool:
                pool.submit(lambda: ProxyConfig.objects.filter(domain_name="zulu.find.example.org").update(
                    domain_name="fresh.example.org"
                )).result()
            expect(page.locator("#proxies .search-match")).to_have_count(1, timeout=10000)
            expect(page.locator("#proxies .search-match")).to_contain_text("fresh.example.org")
            expect(search).to_have_value("fresh.example.org")
            expect(status).to_contain_text("1 matching item")
            expect(page.locator("#proxies .search-blink")).to_have_count(0)
            search.fill("")
            expect(page.locator(".search-match")).to_have_count(0)
            self.assertEqual(page.evaluate("window.documentToken"), "unchanged")
            self.assertEqual(errors, [])
            browser.close()

    def test_account_menu_updates_contacts_and_password_without_navigation(self):
        get_user_model().objects.create_superuser("root", password="original-test-password")
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(channel="chrome", headless=True)
            page = browser.new_page()
            errors = []
            page.on("pageerror", lambda error: errors.append(str(error)))
            page.goto(self.live_server_url + "/login/")
            page.evaluate("window.documentToken = 'unchanged'")
            page.locator('[name="username"]').fill("root")
            page.locator('[name="password"]').fill("original-test-password")
            page.get_by_role("button", name="Sign in", exact=True).click()
            expect(page.locator(".stats")).to_be_visible()

            menu = page.locator(".account-dropdown")
            expect(menu.locator("summary")).to_have_text("R")
            menu.locator("summary").click()
            for label in ["Account information", "Change password", "Add or change mobile number", "Add or change email ID"]:
                expect(menu.get_by_role("link", name=label, exact=True)).to_be_visible()
            expect(menu.get_by_role("button", name="Logout", exact=True)).to_be_visible()
            page.keyboard.press("Escape")
            expect(menu).not_to_have_attribute("open", "")
            menu.locator("summary").click()
            page.locator("h1").click()
            expect(menu).not_to_have_attribute("open", "")
            menu.locator("summary").click()
            menu.get_by_role("link", name="Account information", exact=True).click()
            expect(page.locator("h1")).to_have_text("Account information")
            expect(page.locator(".account-info")).to_contain_text("root")

            menu.locator("summary").click()
            menu.get_by_role("link", name="Add or change mobile number", exact=True).click()
            page.get_by_label("Mobile number", exact=True).fill("+91 9876543210")
            page.get_by_label("Email ID", exact=True).fill("root@example.org")
            page.get_by_role("button", name="Save changes", exact=True).click()
            expect(page.locator(".account-info")).to_contain_text("+91 9876543210")
            expect(page.locator(".account-info")).to_contain_text("root@example.org")

            menu.locator("summary").click()
            menu.get_by_role("link", name="Add or change email ID", exact=True).click()
            expect(page.get_by_label("Mobile number", exact=True)).to_have_value("+91 9876543210")
            expect(page.get_by_label("Email ID", exact=True)).to_have_value("root@example.org")
            page.get_by_label("Email ID", exact=True).fill("updated@example.org")
            page.get_by_role("button", name="Save changes", exact=True).click()
            expect(page.locator(".account-info")).to_contain_text("updated@example.org")

            menu.locator("summary").click()
            menu.get_by_role("link", name="Change password", exact=True).click()
            page.locator('[name="old_password"]').fill("original-test-password")
            page.locator('[name="new_password1"]').fill("changed-test-password-2026!")
            page.locator('[name="new_password2"]').fill("changed-test-password-2026!")
            page.get_by_role("button", name="Save changes", exact=True).click()
            expect(page.locator(".messages")).to_contain_text("Password changed successfully.")
            expect(page.locator(".account-info")).to_contain_text("updated@example.org")
            menu.locator("summary").click()
            menu.get_by_role("button", name="Logout", exact=True).click()
            expect(page.locator(".login-panel")).to_be_visible()

            page.locator('[name="username"]').fill("root")
            page.locator('[name="password"]').fill("changed-test-password-2026!")
            page.get_by_role("button", name="Sign in", exact=True).click()
            expect(page.locator(".stats")).to_be_visible()
            page.set_viewport_size({"width": 390, "height": 844})
            expect(menu.locator("summary")).to_be_visible()
            menu.locator("summary").click()
            expect(menu.get_by_role("link", name="Account information", exact=True)).to_be_visible()
            page.keyboard.press("Escape")
            expect(menu).not_to_have_attribute("open", "")
            self.assertEqual(page.evaluate("window.documentToken"), "unchanged")
            self.assertEqual(errors, [])
            browser.close()
