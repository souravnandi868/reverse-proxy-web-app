from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import TestCase
from django.urls import reverse

from .captive_models import CaptivePortalUser
from io import BytesIO
from openpyxl import load_workbook
from .authorized_users_excel import build_authorized_users_excel


class AuthorizedUsersExcelTests(TestCase):
    def setUp(self):
        self.viewer = get_user_model().objects.create_user("report-viewer", is_staff=True)
        self.viewer.user_permissions.add(Permission.objects.get(codename="view_captiveportaluser"))
        self.url = reverse("captive_users_export_excel")

    def test_access_and_get_only(self):
        self.assertEqual(self.client.get(self.url).status_code, 302)
        self.viewer.user_permissions.clear()
        self.client.force_login(self.viewer)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.viewer.user_permissions.add(Permission.objects.get(codename="view_captiveportaluser"))
        self.assertEqual(self.client.post(self.url).status_code, 405)

    def test_export_all_pages_ignores_filters_and_excludes_deleted_users(self):
        for index in range(27):
            CaptivePortalUser.objects.create(name=f"User {index:02}", section="IT", rank="Officer",
                mobile_number=f"+91900000{index:04}", email_address=f"user{index}@example.org",
                is_enabled=index % 2 == 0, all_fqdns=True)
        deleted = CaptivePortalUser.objects.first()
        deleted.delete()
        self.client.force_login(self.viewer)
        self.assertContains(self.client.get(reverse("captive_users")), self.url)
        with patch("proxies.authorized_users_excel.build_authorized_users_excel", wraps=build_authorized_users_excel) as builder:
            response = self.client.get(self.url, {"page": 2, "q": "No match", "status": "enabled"})
            exported = list(builder.call_args.args[0])
        self.assertEqual(len(exported), 26)
        self.assertNotIn(deleted.pk, [user.pk for user in exported])
        self.assertTrue(any(not user.is_enabled for user in exported))
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn('attachment; filename="authorized-users.xlsx"', response["Content-Disposition"])
        self.assertIn("no-store", response["Cache-Control"])
        sheet = load_workbook(BytesIO(response.content)).active
        data = list(sheet.values)
        self.assertEqual(len(data), 27)
        self.assertEqual(data[0][1:8], ("Name", "Section", "Rank", "Mobile number", "Email address", "Authorized FQDNs", "Status"))
        self.assertEqual(data[1][4], "+919000000001")
        self.assertEqual(data[1][6], "All captive-enabled FQDNs")
        self.assertEqual(data[1][7], "Disabled")
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:M27")

    def test_empty_export(self):
        self.client.force_login(self.viewer)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(load_workbook(BytesIO(response.content)).active.max_row, 1)

    def test_untrusted_names_are_literal_text(self):
        CaptivePortalUser.objects.create(name="=1+1", section="IT", rank="Officer",
            mobile_number="+919000000099", email_address="formula@example.org")
        self.client.force_login(self.viewer)
        sheet = load_workbook(BytesIO(self.client.get(self.url).content)).active
        self.assertEqual(sheet["B2"].value, "=1+1")
        self.assertEqual(sheet["B2"].data_type, "s")
