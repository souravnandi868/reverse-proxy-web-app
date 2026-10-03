"""Tabular authorized-user workbook with literal, wrapping text values."""
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def build_authorized_users_excel(users):
    workbook = Workbook(write_only=True)
    headers = ["ID", "Name", "Section", "Rank", "Mobile number", "Email address",
               "Authorized FQDNs", "Status", "Last successful login", "Created at",
               "Created by", "Updated at", "Updated by"]

    def timestamp(value):
        return timezone.localtime(value).isoformat(sep=" ") if value else "Never"

    def cell(sheet, value):
        result = WriteOnlyCell(sheet, ILLEGAL_CHARACTERS_RE.sub("", str(value if value is not None else "")))
        result.data_type = "s"
        result.alignment = Alignment(vertical="top", wrap_text=True)
        return result

    def new_sheet():
        sheet = workbook.create_sheet(f"Authorized users {len(workbook.worksheets) + 1}")
        sheet.freeze_panes = "A2"
        for index, width in enumerate([10, 30, 25, 25, 20, 35, 55, 14, 32, 32, 25, 32, 25], 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        row = [cell(sheet, value) for value in headers]
        for item in row:
            item.font = Font(bold=True, color="FFFFFF")
            item.fill = PatternFill("solid", fgColor="15345F")
        sheet.append(row)
        return sheet

    sheet = new_sheet()
    count = 1
    for user in users:
        if count == 1048576:
            sheet.auto_filter.ref = f"A1:M{count}"
            sheet = new_sheet()
            count = 1
        domains = "All captive-enabled FQDNs" if user.all_fqdns else ", ".join(
            proxy.domain_name for proxy in user.proxies.all()) or "None"
        values = [user.pk, user.name, user.section, user.rank, user.mobile_number, user.email_address,
                  domains, "Enabled" if user.is_enabled else "Disabled", timestamp(user.last_successful_login),
                  timestamp(user.created_at), user.created_by or "Not set",
                  timestamp(user.updated_at), user.updated_by or "Not set"]
        sheet.append([cell(sheet, value) for value in values])
        count += 1
    sheet.auto_filter.ref = f"A1:M{count}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
