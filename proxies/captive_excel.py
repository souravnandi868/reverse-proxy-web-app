"""Captive audit export; untrusted values are stored as literal Excel text."""
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill


def build_captive_audit_excel(events):
    workbook = Workbook(write_only=True)
    headers = [f"Time ({timezone.get_current_timezone_name()})", "Administrator", "User",
               "FQDN", "Action", "Channel", "Client IP"]

    def cell(sheet, value):
        result = WriteOnlyCell(sheet, value=ILLEGAL_CHARACTERS_RE.sub("", str(value or "")))
        result.data_type = "s"
        return result

    def new_sheet():
        sheet = workbook.create_sheet(f"Captive audit {len(workbook.worksheets) + 1}")
        sheet.freeze_panes = "A2"
        for index, width in enumerate([30, 25, 30, 40, 30, 15, 25]):
            sheet.column_dimensions[chr(65 + index)].width = width
        row = [cell(sheet, value) for value in headers]
        for item in row:
            item.font = Font(bold=True, color="FFFFFF")
            item.fill = PatternFill("solid", fgColor="15345F")
        sheet.append(row)
        return sheet

    sheet = new_sheet()
    count = 1
    for event in events:
        if count == 1048576:
            sheet.auto_filter.ref = f"A1:G{count}"
            sheet = new_sheet()
            count = 1
        values = [timezone.localtime(event.created_at).isoformat(sep=" "),
                  event.actor, event.user, event.proxy, event.action, event.channel, event.client_ip]
        sheet.append([cell(sheet, value) for value in values])
        count += 1
    sheet.auto_filter.ref = f"A1:G{count}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
