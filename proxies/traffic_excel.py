"""Excel traffic reports. Log text is always stored as text, never as formulas."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill


def build_traffic_excel(entries):
    workbook = Workbook(write_only=True)
    fields = [
        ("Time", "time", 28), ("Incoming IP", "source_ip", 22),
        ("Destination FQDN", "destination_fqdn", 35),
        ("Destination server", "destination_server", 24),
        ("Request", "request", 55), ("Status", "status", 12),
        ("Request time (seconds)", "request_time", 24),
        ("Bytes", "bytes", 16), ("User agent", "user_agent", 55),
    ]

    def text_cell(sheet, value):
        cell = WriteOnlyCell(sheet, value=ILLEGAL_CHARACTERS_RE.sub("", str(value if value is not None else "")))
        cell.data_type = "s"
        return cell

    def new_sheet():
        sheet = workbook.create_sheet(f"Traffic {len(workbook.worksheets) + 1}")
        sheet.freeze_panes = "A2"
        for index, (_, _, width) in enumerate(fields):
            sheet.column_dimensions[chr(65 + index)].width = width
        header = [text_cell(sheet, label) for label, _, _ in fields]
        for cell in header:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="15345F")
        sheet.append(header)
        return sheet

    sheet = new_sheet()
    row_count = 1
    for entry in entries:
        if row_count == 1048576:
            sheet.auto_filter.ref = f"A1:I{row_count}"
            sheet = new_sheet()
            row_count = 1
        sheet.append([text_cell(sheet, entry.get(key, "")) for _, key, _ in fields])
        row_count += 1
    sheet.auto_filter.ref = f"A1:I{row_count}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
